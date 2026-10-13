"""
多格式文件分析
支持：pdf / docx / doc / pptx / ppt / txt / md / json / csv / jpg / png / chm
"""
import os
import json
import csv
import subprocess
import tempfile
import zipfile
from io import StringIO
from pypdf import PdfReader
from docx import Document as DocxDocument
from config import ds_client

# ==================== 常见路径 ====================
SOFFICE_CANDIDATES = [
    r"D:\Software\LibreOffice_25.8.7\program\soffice.exe",
    r"C:\Program Files\LibreOffice\program\soffice.exe",
    r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    "soffice",
]

SEVEN_ZIP_CANDIDATES = [
    r"C:\Program Files\7-Zip\7z.exe",
    r"C:\Program Files (x86)\7-Zip\7z.exe",
    "7z",
]


def _find_exe(candidates) -> str:
    for c in candidates:
        if c == os.path.basename(c):   # 无路径，从 PATH 找
            return c
        if os.path.exists(c):
            return c
    return ""


# ==================== 各格式 → 文本 ====================

def _pdf_to_text(file_bytes):
    """PDF 提取（文本型直接提取，扫描件返回空）"""
    with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
        tmp.write(file_bytes); tmp_path = tmp.name
    try:
        reader = PdfReader(tmp_path)
        parts = []
        for i, page in enumerate(reader.pages):
            t = page.extract_text() or ""
            if t.strip():
                parts.append(f"--- 第 {i+1} 页 ---\n{t}")
        return "\n\n".join(parts)
    finally:
        if os.path.exists(tmp_path): os.remove(tmp_path)


def _docx_to_text(file_bytes):
    with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
        tmp.write(file_bytes); tmp_path = tmp.name
    try:
        doc = DocxDocument(tmp_path)
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(parts)
    finally:
        if os.path.exists(tmp_path): os.remove(tmp_path)


def _doc_to_text(file_bytes):
    """老版 .doc：用 LibreOffice 转 docx 再读"""
    soffice = _find_exe(SOFFICE_CANDIDATES)
    if not soffice:
        return ""
    with tempfile.NamedTemporaryFile(delete=False, suffix=".doc") as tmp:
        tmp.write(file_bytes); tmp_path = tmp.name
    try:
        out_dir = os.path.dirname(tmp_path)
        subprocess.run(
            [soffice, "--headless", "--convert-to", "docx",
             "--outdir", out_dir, tmp_path],
            check=True, capture_output=True, timeout=120,
        )
        docx_path = tmp_path.replace(".doc", ".docx")
        if os.path.exists(docx_path):
            with open(docx_path, "rb") as f:
                return _docx_to_text(f.read())
        return ""
    except Exception:
        return ""
    finally:
        if os.path.exists(tmp_path): os.remove(tmp_path)


def _pptx_to_text(file_bytes):
    """pptx：用 python-pptx 提取每页文本"""
    try:
        from pptx import Presentation
    except ImportError:
        return ""
    with tempfile.NamedTemporaryFile(delete=False, suffix=".pptx") as tmp:
        tmp.write(file_bytes); tmp_path = tmp.name
    try:
        prs = Presentation(tmp_path)
        parts = []
        for i, slide in enumerate(prs.slides, 1):
            texts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    texts.append(shape.text.strip())
            if texts:
                parts.append(f"--- 第 {i} 页 ---\n" + "\n".join(texts))
        return "\n\n".join(parts)
    finally:
        if os.path.exists(tmp_path): os.remove(tmp_path)


def _ppt_to_text(file_bytes):
    """老版 .ppt：用 LibreOffice 转 pptx"""
    soffice = _find_exe(SOFFICE_CANDIDATES)
    if not soffice:
        return ""
    with tempfile.NamedTemporaryFile(delete=False, suffix=".ppt") as tmp:
        tmp.write(file_bytes); tmp_path = tmp.name
    try:
        out_dir = os.path.dirname(tmp_path)
        subprocess.run(
            [soffice, "--headless", "--convert-to", "pptx",
             "--outdir", out_dir, tmp_path],
            check=True, capture_output=True, timeout=120,
        )
        pptx_path = tmp_path.replace(".ppt", ".pptx")
        if os.path.exists(pptx_path):
            with open(pptx_path, "rb") as f:
                return _pptx_to_text(f.read())
        return ""
    except Exception:
        return ""
    finally:
        if os.path.exists(tmp_path): os.remove(tmp_path)


def _image_to_text(file_bytes, file_name):
    """jpg/png：OCR（需要 pytesseract + tesseract）"""
    try:
        import pytesseract
        from PIL import Image
    except ImportError:
        return ""
    with tempfile.NamedTemporaryFile(delete=False,
                                     suffix=os.path.splitext(file_name)[1]) as tmp:
        tmp.write(file_bytes); tmp_path = tmp.name
    try:
        img = Image.open(tmp_path)
        text = pytesseract.image_to_string(img, lang="chi_sim")
        return text
    except Exception:
        return ""
    finally:
        if os.path.exists(tmp_path): os.remove(tmp_path)


def _chm_to_text(file_bytes):
    """
    chm：用 7z 解压出 HTML，再把 HTML 转文本
    7z 命令行：7z x input.chm -o<output_dir>
    """
    seven_zip = _find_exe(SEVEN_ZIP_CANDIDATES)
    if not seven_zip:
        return ""
    with tempfile.TemporaryDirectory() as tmp_dir:
        chm_path = os.path.join(tmp_dir, "input.chm")
        with open(chm_path, "wb") as f:
            f.write(file_bytes)
        out_dir = os.path.join(tmp_dir, "extracted")
        try:
            subprocess.run(
                [seven_zip, "x", chm_path, f"-o{out_dir}", "-y"],
                check=True, capture_output=True, timeout=300,
            )
        except Exception:
            return ""

        # 遍历解压出的 HTML / txt / htm
        texts = []
        for root, _, files in os.walk(out_dir):
            for fname in files:
                if fname.lower().endswith((".htm", ".html", ".txt")):
                    try:
                        with open(os.path.join(root, fname), "r",
                                  encoding="utf-8", errors="ignore") as f:
                            texts.append(f.read())
                    except Exception:
                        pass
        # 简单去 HTML 标签
        import re
        combined = "\n\n".join(texts)
        combined = re.sub(r"<script[^>]*>.*?</script>", "", combined, flags=re.DOTALL)
        combined = re.sub(r"<style[^>]*>.*?</style>", "", combined, flags=re.DOTALL)
        combined = re.sub(r"<[^>]+>", " ", combined)
        combined = re.sub(r"\s+", " ", combined)
        return combined


def _txt_to_text(file_bytes):
    for enc in ["utf-8", "gbk", "gb18030"]:
        try:
            return file_bytes.decode(enc)
        except UnicodeDecodeError:
            continue
    return file_bytes.decode("utf-8", errors="ignore")


def _json_to_text(file_bytes):
    try:
        data = json.loads(file_bytes.decode("utf-8"))
        return json.dumps(data, ensure_ascii=False, indent=2)
    except Exception:
        return ""


def _csv_to_text(file_bytes):
    text = _txt_to_text(file_bytes)
    try:
        rows = list(csv.reader(StringIO(text)))
        return "\n".join(" | ".join(row) for row in rows)
    except Exception:
        return text


# ==================== 统一入口 ====================

def extract_text(file_bytes: bytes, file_name: str) -> str:
    """根据扩展名选择提取方式"""
    ext = file_name.lower().rsplit(".", 1)[-1] if "." in file_name else ""
    if ext == "pdf":   return _pdf_to_text(file_bytes)
    if ext == "docx":  return _docx_to_text(file_bytes)
    if ext == "doc":   return _doc_to_text(file_bytes)
    if ext == "pptx":  return _pptx_to_text(file_bytes)
    if ext == "ppt":   return _ppt_to_text(file_bytes)
    if ext in ("jpg", "jpeg", "png"): return _image_to_text(file_bytes, file_name)
    if ext == "chm":   return _chm_to_text(file_bytes)
    if ext in ("txt", "md"): return _txt_to_text(file_bytes)
    if ext == "json":  return _json_to_text(file_bytes)
    if ext == "csv":   return _csv_to_text(file_bytes)
    return _txt_to_text(file_bytes)


# ==================== LLM 抽取案情 ====================

CASE_PROMPT = """你是一位消防行政执法人员。请从以下文档中提取案情信息。

文档内容：
{text}

请严格返回 JSON（缺失字段填 null，不要编造）：
{{
  "case_context": "详细描述违法事实（时间、地点、当事人、违法行为）",
  "unit_name": "涉案单位全称",
  "unit_legal_person": "法定代表人姓名",
  "unit_address": "单位地址",
  "unit_contact": "联系电话",
  "unit_credit_code": "统一社会信用代码",
  "contact_person": "联系人姓名",
  "contact_phone": "联系人电话",
  "interviewee": "被询问人（姓名 + 职务）",
  "time": "违法时间 YYYY年M月D日",
  "location": "违法地点"
}}
"""


def analyze_case_file(file_bytes: bytes, file_name: str) -> dict:
    """从任意格式文件中抽取案情"""
    text = extract_text(file_bytes, file_name)
    if not text.strip():
        return {"_error": f"无法从 {file_name} 提取文本（可能是扫描件或空文件）"}
    text = text[:8000]

    resp = ds_client.chat.completions.create(
        model="deepseek-chat",
        messages=[{"role": "user", "content": CASE_PROMPT.format(text=text)}],
        temperature=0.1,
        response_format={"type": "json_object"},
    )
    raw = json.loads(resp.choices[0].message.content)
    return {k: v for k, v in raw.items() if v}