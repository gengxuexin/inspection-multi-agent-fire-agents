# render_from_facts.py
import os
from docx import Document
from document_io import fill_paragraph_placeholders, fill_table_cells


def render_docx_from_facts(
    template_path: str,
    output_path: str,
    registry,
    include_assumptions: bool = True,
    include_manual: bool = False,
):
    """
    用事实库渲染 docx。
    include_assumptions=True 时，测试假设字段会填入；
    include_manual=False 时，必须人工确认字段留空。
    """
    mapping = registry.to_mapping(
        include_assumptions=include_assumptions,
        include_manual=include_manual,
    )

    doc = Document(template_path)
    fill_paragraph_placeholders(doc, mapping)
    for table in doc.tables:
        fill_table_cells(table, mapping)

    doc.save(output_path)
    return output_path, registry.missing_required()