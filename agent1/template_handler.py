"""
模板处理：
1. 检测模板是否已加占位符（【xxx】/{{xxx}}）
2. 提取模板中所有占位符名字
3. 若传入 facts_registry，则直接用事实库渲染；
   否则走原来的 LLM 智能映射逻辑
"""
import os
import json
import re
from docx import Document
from config import ds_client
from document_io import (
    fill_table_cells, fill_paragraph_placeholders,
    PLACEHOLDER_PATTERN,
)


def has_placeholders(docx_path: str) -> bool:
    doc = Document(docx_path)
    text = "\n".join(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text += "\n" + cell.text
    return bool(PLACEHOLDER_PATTERN.search(text))


def extract_placeholders(docx_path: str) -> list:
    doc = Document(docx_path)
    text = "\n".join(p.text for p in doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text += "\n" + cell.text

    seen, result = set(), []
    for m in PLACEHOLDER_PATTERN.finditer(text):
        name = m.group(1) or m.group(2)
        name = name.replace(" ", "").replace("　", "")
        if name and name not in seen:
            seen.add(name)
            result.append(name)
    return result


def extract_labels(docx_path: str) -> list:
    doc = Document(docx_path)
    labels = []

    for para in doc.paragraphs:
        t = para.text.strip()
        m = re.match(r"^([\u4e00-\u9fa5A-Za-z0-9_（）()]{2,20})[：:]\s*[_\-—]{2,}", t)
        if m:
            labels.append(m.group(1))

    for table in doc.tables:
        for row in table.rows:
            cells = row.cells
            for i, cell in enumerate(cells):
                t = cell.text.strip().rstrip("：:")
                if not (2 <= len(t) <= 20):
                    continue
                if PLACEHOLDER_PATTERN.search(t):
                    continue
                if i + 1 < len(cells):
                    right = cells[i + 1].text.strip()
                    if not right or PLACEHOLDER_PATTERN.search(right):
                        labels.append(t)

    seen, result = set(), []
    for lb in labels:
        if lb not in seen:
            seen.add(lb)
            result.append(lb)
    return result


# ==================== LLM 智能映射（仅 fallback） ====================

MAP_PROMPT = """你是一位消防行政执法人员。
请根据以下案情信息，为文书模板中的每个占位符填写合适的值。

案情信息：
{case_info}

文书占位符列表（请逐个填写）：
{placeholders}

要求：
1. 严格返回 JSON，key 就是占位符名称（不要加【】）
2. 能为每个占位符找到或推断出合适的值就填，找不到填 null
3. 常见占位符的默认值参考：
   - 年份：从案情时间提取，如 2024
   - 文号：6 位数字，如 000052
   - 日期/决定日期/检查日期：案情时间
   - 消防救援机构名称：如 桐乡市消防救援大队
   - 复议机关：如 桐乡市人民政府
   - 管辖法院：如 桐乡市人民
   - 执行方式和期限：自收到本决定书之日起十五日内到指定银行缴纳罚款
   - 证据：现场检查记录、询问笔录、现场照片
   - 违法行为人信息：单位名称/地址/法定代表人/统一社会信用代码
4. 具体返回示例：{{"年份": "2024", "文号": "000052", "案由": "..."}}

请只返回 JSON，不要任何解释。
"""


def _llm_map(placeholders: list, case_context: dict) -> dict:
    case_info = json.dumps(case_context, ensure_ascii=False, indent=2)
    pl_json = json.dumps(placeholders, ensure_ascii=False)

    resp = ds_client.chat.completions.create(
        model="deepseek-chat",
        messages=[{"role": "user", "content": MAP_PROMPT.format(
            case_info=case_info,
            placeholders=pl_json,
        )}],
        temperature=0.1,
        response_format={"type": "json_object"},
    )
    raw = json.loads(resp.choices[0].message.content)
    return {k: v for k, v in raw.items() if v is not None and v != ""}


# ==================== 回填 ====================

def render_template(docx_path, output_path, case_context,
                    facts_registry=None,
                    include_assumptions=True,
                    include_manual=False):
    """
    facts_registry: FactRegistry 实例；非 None 时使用事实库回填，
                    否则走原来的 LLM 智能映射逻辑。
    """
    doc = Document(docx_path)

    # ---------- 事实库模式 ----------
    if facts_registry is not None:
        if hasattr(facts_registry, "to_mapping"):
            mapping = facts_registry.to_mapping(
                include_assumptions=include_assumptions,
                include_manual=include_manual,
            )
        else:
            # 兼容直接传 dict
            mapping = {k: ("" if v is None else str(v))
                       for k, v in dict(facts_registry).items()}

        fill_paragraph_placeholders(doc, mapping)
        for table in doc.tables:
            fill_table_cells(table, mapping)

        doc.save(output_path)
        return output_path

    # ---------- 旧逻辑：占位符模式 ----------
    if has_placeholders(docx_path):
        placeholders = extract_placeholders(docx_path)
        mapping = {k: ("" if v is None else str(v))
                   for k, v in case_context.items()}
        try:
            llm_mapping = _llm_map(placeholders, case_context)
            mapping.update(llm_mapping)
        except Exception as e:
            print(f"[warn] LLM 映射失败: {e}")

        fill_paragraph_placeholders(doc, mapping)
        for table in doc.tables:
            fill_table_cells(table, mapping)

    # ---------- 旧逻辑：标签模式 ----------
    else:
        labels = extract_labels(docx_path)
        if not labels:
            raise ValueError("模板中未识别到字段，请用【占位符】标注或提供带标签的表格")

        mapping = _llm_map(labels, case_context)
        _smart_fill_paragraphs(doc, mapping)
        _smart_fill_tables(doc, mapping)

    doc.save(output_path)
    return output_path


def _smart_fill_paragraphs(doc, mapping: dict):
    for para in doc.paragraphs:
        t = para.text
        for label, value in mapping.items():
            pattern = rf"({re.escape(label)}[：:]\s*)[_\-—]{{2,}}"
            if re.search(pattern, t):
                t = re.sub(pattern, rf"\g<1>{value}", t)
        if t != para.text and para.runs:
            para.runs[0].text = t
            for r in para.runs[1:]:
                r.text = ""


def _smart_fill_tables(doc, mapping: dict):
    for table in doc.tables:
        label_index = {}
        for i, row in enumerate(table.rows):
            for j, cell in enumerate(row.cells):
                key = cell.text.strip().rstrip("：:")
                if key:
                    label_index[key] = (i, j)

        for label, value in mapping.items():
            if label not in label_index:
                continue
            i, j = label_index[label]
            target = None
            if j + 1 < len(table.rows[i].cells):
                right = table.rows[i].cells[j + 1]
                if right.text.strip() == "":
                    target = right
            if target is None and i + 1 < len(table.rows):
                down = table.rows[i + 1].cells[j]
                if down.text.strip() == "":
                    target = down
            if target is not None:
                if target.paragraphs and target.paragraphs[0].runs:
                    target.paragraphs[0].runs[0].text = str(value)
                    for r in target.paragraphs[0].runs[1:]:
                        r.text = ""
                else:
                    target.text = str(value)