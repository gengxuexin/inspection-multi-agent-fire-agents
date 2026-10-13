import os
import re
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
import pdfplumber


# 占位符正则：{{key}} 或 【key】
PLACEHOLDER_PATTERN = re.compile(r"\{\{(\w+)\}\}|【([^】]+)】")


def _normalize_nested(text: str) -> str:
    """把 【【年份】】 这种双层括号压成 【年份】"""
    prev = None
    while prev != text:
        prev = text
        text = re.sub(r"【【([^【】]+)】】", r"【\1】", text)
    return text


def _replace_in_text(full_text: str, norm_mapping: dict) -> str:
    full_text = _normalize_nested(full_text)

    def replace_match(m):
        key = m.group(1) or m.group(2)
        key_norm = key.replace(" ", "").replace("　", "")
        if key_norm in norm_mapping:
            v = norm_mapping[key_norm]
            return "" if v is None else str(v)
        return m.group(0)

    return PLACEHOLDER_PATTERN.sub(replace_match, full_text)


# ==================== Word 读取 ====================
def read_docx(path: str) -> dict:
    doc = Document(path)
    paragraphs = [p.text for p in doc.paragraphs]
    tables = []
    for table in doc.tables:
        table_data = []
        for row in table.rows:
            row_data = [cell.text for cell in row.cells]
            table_data.append(row_data)
        tables.append(table_data)
    return {"paragraphs": paragraphs, "tables": tables, "doc_object": doc}


# ==================== 段落回填 ====================
def fill_paragraph_placeholders(doc, mapping: dict):
    """将 {{key}} / 【key】 / 【【key】】替换成 mapping[key]，保留原格式"""
    norm_mapping = {k.replace(" ", "").replace("　", ""): v
                    for k, v in mapping.items()}

    for para in doc.paragraphs:
        full_text = "".join(run.text for run in para.runs)
        if not full_text:
            continue
        new_text = _replace_in_text(full_text, norm_mapping)
        if new_text == full_text or not para.runs:
            continue
        para.runs[0].text = new_text
        for run in para.runs[1:]:
            run.text = ""


# ==================== 表格单元格回填 ====================
def _write_cell(cell, value: str):
    if cell.paragraphs and cell.paragraphs[0].runs:
        cell.paragraphs[0].runs[0].text = str(value)
        for r in cell.paragraphs[0].runs[1:]:
            r.text = ""
    else:
        cell.text = str(value)


def fill_table_cells(table, cell_mapping: dict):
    """
    1. 表格里所有单元格内部的占位符直接替换（含【【】】双层）
    2. "标签格 → 右侧/下方空单元格"的传统模式也回填
    """
    norm_mapping = {k.replace(" ", "").replace("　", ""): v
                    for k, v in cell_mapping.items()}

    # 1. 单元格内占位符
    for row in table.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                full_text = "".join(run.text for run in para.runs)
                if not full_text:
                    continue
                new_text = _replace_in_text(full_text, norm_mapping)
                if new_text != full_text and para.runs:
                    para.runs[0].text = new_text
                    for r in para.runs[1:]:
                        r.text = ""

    # 2. 标签 + 右/下空格
    label_index = {}
    for i, row in enumerate(table.rows):
        for j, cell in enumerate(row.cells):
            text = cell.text.strip().rstrip("：:").replace(" ", "").replace("　", "")
            if text:
                label_index[text] = (i, j)

    for label, value in cell_mapping.items():
        if not value:
            continue
        key = label.replace(" ", "").replace("　", "")
        if key not in label_index:
            continue
        i, j = label_index[key]
        target = None
        if j + 1 < len(table.rows[i].cells):
            right = table.rows[i].cells[j + 1]
            if right.text.strip() == "":
                target = right
        if target is None and i + 1 < len(table.rows):
            down = table.rows[i + 1].cells[j]
            if down.text.strip() == "":
                target = down
        if target is not None:
            _write_cell(target, str(value))


# ==================== PDF 读取 ====================
def read_pdf_tables(path: str) -> list:
    all_tables = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            for t in page.extract_tables():
                all_tables.append(t)
    return all_tables


def read_pdf_text(path: str) -> str:
    text = ""
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            text += (page.extract_text() or "") + "\n"
    return text