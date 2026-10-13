"""
消防法律智能助手 - 统一入口
文件分析 · 模板驱动文书生成（含 RAG 法律推理）· RAG 法律问答
"""
import os
import sys
import time
import hashlib
import tempfile
import zipfile
from datetime import datetime
from io import BytesIO
from field_resolver import FieldKind
import streamlit as st
from docx import Document

_SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
if _SCRIPT_DIR not in sys.path:
    sys.path.insert(0, _SCRIPT_DIR)

from config import set_deepseek, set_embedding, get_deepseek, get_embedding_config
from document_io import fill_table_cells, fill_paragraph_placeholders
from file_analyzer import analyze_case_file
from pdf_converter import docx_to_pdf
from template_schema_extractor import extract_fields_from_template, summarize_fields
from case_field_matcher import (
    extract_fields_from_case, verify_extraction, build_review_table,
)
from rag_legal_advisor import split_fields, is_legal_field


# ==================== 页面配置 ====================
st.set_page_config(
    page_title="消防法律智能助手",
    page_icon="⚖️",
    layout="wide",
    initial_sidebar_state="expanded",
)

st.markdown("""
<style>
.block-container {padding-top: 2rem; padding-bottom: 3rem;}
h1 {font-weight: 600; letter-spacing: -0.5px;}
.stTabs [data-baseweb="tab"] {padding: 12px 24px; font-weight: 500;}
.stButton>button {border-radius: 6px; font-weight: 500;}
</style>
""", unsafe_allow_html=True)


# ==================== 常量 ====================
OUTPUT_DIR = os.path.join(_SCRIPT_DIR, "output")
INDEX_DIRS = {
    "single": os.path.join(_SCRIPT_DIR, "faiss_single"),
    "multi":  os.path.join(_SCRIPT_DIR, "faiss_multi"),
    "full":   os.path.join(_SCRIPT_DIR, "faiss_full"),
}


# ==================== session state ====================
if "case_dict" not in st.session_state:
    st.session_state.case_dict = {}


# ==================== 侧边栏 ====================
with st.sidebar:
    st.markdown("### API 配置")
    ds_key = st.text_input("DeepSeek API Key", type="password", placeholder="sk-...")
    emb_key = st.text_input("Embedding API Key", type="password", placeholder="sk-...")

    st.divider()
    st.markdown("### 检索参数")
    chunk_size    = st.slider("Chunk Size",    100, 2000, 500, 50)
    chunk_overlap = st.slider("Chunk Overlap",   0,  500, 100, 20)
    top_k         = st.slider("Top-K",           1,   10,   4,  1)

    st.divider()
    st.markdown("### 文书生成设置")
    use_rag_legal = st.checkbox(
        "启用 RAG 法律推理",
        value=True,
        help="法律依据/违法事实/触犯条款类字段，从本地消防法律知识库检索，而不是仅靠案情描述",
    )
    legal_top_k = st.slider(
        "RAG 检索片段数", 3, 15, 8, 1,
        help="检索多少个法条片段供 LLM 归纳（越大越全，也越慢）",
    )

    st.divider()
    if ds_key:
        set_deepseek(ds_key.strip())
        st.success("DeepSeek ✓", icon="✅")
    else:
        st.warning("请先填写 DeepSeek API Key")
    if emb_key:
        set_embedding(emb_key.strip())
        st.success("Embedding ✓", icon="✅")
    else:
        st.warning("请先填写 Embedding API Key")

    key_sig = f"{ds_key.strip()[:15]}_{emb_key.strip()[:15]}"
    if st.session_state.get("_key_sig") != key_sig:
        st.session_state["_key_sig"] = key_sig
        for k in list(st.session_state.keys()):
            if k.startswith("rag_"):
                del st.session_state[k]
        st.info("检测到 API Key 变化，已重置 RAG 缓存")


# ==================== 顶部标题 ====================
st.title("消防法律智能助手")
st.caption("文件分析 · 模板驱动文书生成 · RAG 法律问答")


tab_analyze, tab_generate, tab_qa = st.tabs(["文件分析", "文书生成", "法律问答"])


# ============================================================
# Tab1: 文件分析
# ============================================================
with tab_analyze:
    st.markdown("#### 上传案情材料，自动识别关键信息")
    st.caption("支持 PDF / DOCX / DOC / TXT / MD / JSON / CSV")

    up = st.file_uploader(
        "选择文件",
        type=["pdf", "docx", "doc", "txt", "md", "json", "csv"],
        key="file_analyze",
        label_visibility="collapsed",
    )

    col1, col2 = st.columns([1, 4])
    if up and col1.button("开始分析", type="primary", use_container_width=True):
        with st.spinner("正在分析文件..."):
            try:
                info = analyze_case_file(up.getvalue(), up.name)
            except Exception as e:
                st.error(f"分析失败: {e}")
                info = {}
        if info:
            if "_error" in info:
                st.error(info["_error"])
            else:
                st.success("分析完成，可在下方编辑后应用到案情")
                st.session_state.case_dict.update({k: v for k, v in info.items() if v})

    if st.session_state.case_dict:
        st.divider()
        st.markdown("##### 当前案情")
        with st.form("case_edit_form"):
            c1, c2, c3 = st.columns(3)
            case_id = c1.text_input("案件编号", st.session_state.case_dict.get("case_id", ""))
            c_time  = c2.text_input("时间",     st.session_state.case_dict.get("time", ""))
            c_loc   = c3.text_input("地点",     st.session_state.case_dict.get("location", ""))

            c1, c2, c3 = st.columns(3)
            interviewer = c1.text_input("承办人",   st.session_state.case_dict.get("interviewer", ""))
            interviewee = c2.text_input("被询问人", st.session_state.case_dict.get("interviewee", ""))
            leader      = c3.text_input("负责人",   st.session_state.case_dict.get("负责人", ""))

            st.markdown("**单位信息**")
            c1, c2 = st.columns(2)
            u_name  = c1.text_input("单位名称",   st.session_state.case_dict.get("unit_name", ""))
            u_legal = c2.text_input("法定代表人", st.session_state.case_dict.get("unit_legal_person", ""))
            u_addr  = c1.text_input("单位地址",   st.session_state.case_dict.get("unit_address", ""))
            u_tel   = c2.text_input("联系电话",   st.session_state.case_dict.get("unit_contact", ""))
            u_code  = c1.text_input("统一社会信用代码", st.session_state.case_dict.get("unit_credit_code", ""))
            c_person = c2.text_input("联系人",   st.session_state.case_dict.get("contact_person", ""))
            c_phone  = c1.text_input("联系人电话", st.session_state.case_dict.get("contact_phone", ""))

            st.markdown("**案情描述**")
            desc = st.text_area("案情", st.session_state.case_dict.get("case_context", ""), height=140)

            if st.form_submit_button("保存案情", type="primary"):
                st.session_state.case_dict.update({
                    "case_id": case_id, "time": c_time, "location": c_loc,
                    "interviewer": interviewer, "interviewee": interviewee,
                    "负责人": leader,
                    "unit_name": u_name, "unit_legal_person": u_legal, "legal_person": u_legal,
                    "unit_address": u_addr, "unit_contact": u_tel,
                    "unit_credit_code": u_code,
                    "contact_person": c_person, "contact_phone": c_phone,
                    "case_context": desc,
                })
                st.success("案情已保存，可切换到「文书生成」")


# ============================================================
# Tab2: 文书生成（模板驱动 + RAG 法律推理）
# ============================================================
with tab_generate:
    st.markdown("#### 上传文书模板，从案情和本地知识库自动抽取字段并回填")

    if not st.session_state.case_dict.get("case_context"):
        st.info("请先在「文件分析」中上传案情，或手动填写案情描述")
    else:
        # ========== Step 1: 上传模板 ==========
        st.markdown("##### 1. 上传模板")
        st.caption("字段来源：占位符【xxx】、标签'XX：___'、表格标签格")

        tpls = st.file_uploader(
            "上传模板",
            type=["docx"],
            accept_multiple_files=True,
            key="template_upload",
            label_visibility="collapsed",
        )

        tpl_sig = tuple(sorted(t.name for t in tpls)) if tpls else ()
        if st.session_state.get("_tpl_sig") != tpl_sig:
            st.session_state["_tpl_sig"] = tpl_sig
            for k in ["_tpl_fields", "_review_rows", "_legal_keys"]:
                st.session_state.pop(k, None)

        if tpls:
            all_tpl_fields = {}
            for tpl in tpls:
                with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
                    tmp.write(tpl.getvalue()); tpl_path = tmp.name
                try:
                    fields = extract_fields_from_template(tpl_path)
                    all_tpl_fields[tpl.name] = fields
                finally:
                    os.remove(tpl_path)

            # 显示统计
            for name, fs in all_tpl_fields.items():
                if not fs:
                    st.warning(f"**{name}** — 未识别到可填字段", icon="⚠️")
                    continue

                sm = summarize_fields(fs)
                normal, legal = split_fields(fs)
                st.success(
                    f"**{name}** — 共 {len(fs)} 个字段 "
                    f"（普通 {len(normal)}，法律 {len(legal)}）"
                    f"  占位符 {sm['placeholder']} / 标签 {sm['label']} / 表格 {sm['table_label']}",
                    icon="✅",
                )
                if legal and use_rag_legal:
                    st.caption("🔍 法律字段将由 RAG 知识库生成：" +
                               "、".join(f["key"] for f in legal))
                if legal and not use_rag_legal:
                    st.warning("⚠️ 检测到法律字段，但侧边栏已关闭 RAG 法律推理")

            # ========== Step 2: 匹配字段 ==========
            st.divider()
            st.markdown("##### 2. 匹配字段")
            if use_rag_legal:
                st.caption("普通字段从案情抽取；法律字段从本地 RAG 知识库检索并归纳")
            else:
                st.caption("仅从案情抽取（RAG 法律推理已关闭）")

            if st.button("开始匹配字段", type="primary", use_container_width=True):
                merged = {}
                for fs in all_tpl_fields.values():
                    for f in fs:
                        merged.setdefault(f["key"], f)
                merged_list = list(merged.values())

                if not merged_list:
                    st.error("没有任何可填字段")
                else:
                    case_text = st.session_state.case_dict["case_context"]

                    with st.spinner(f"抽取 {len(merged_list)} 个字段..."):
                        try:
                            extracted = extract_fields_from_case(
                                case_text,
                                merged_list,
                                case_dict=st.session_state.case_dict,
                                use_rag_for_legal=use_rag_legal,
                            )
                        except Exception as e:
                            st.error(f"抽取失败: {e}")
                            extracted = {}

                    with st.spinner("交叉校验中..."):
                        try:
                            verified = verify_extraction(case_text, extracted)
                        except Exception as e:
                            st.error(f"校验失败: {e}")
                            verified = {}

                    rows = build_review_table(merged_list, extracted, verified)
                    st.session_state["_review_rows"] = rows
                    st.session_state["_tpl_fields"] = all_tpl_fields
                    st.success(f"匹配完成：{len(rows)} 个字段，请审核")

        # ========== Step 3: 字段确认表 ==========
        if st.session_state.get("_review_rows"):
            st.divider()
            st.markdown("##### 3. 字段确认表")
            st.caption("可直接编辑「抽取值」，勾选「确认填入」的字段将回填模板。")

            edited = st.data_editor(
                st.session_state["_review_rows"],
                column_config={
                    "字段":     st.column_config.TextColumn("字段",     disabled=True, width="small"),
                    "来源":     st.column_config.TextColumn("来源",     disabled=True, width="small"),
                    "抽取值":   st.column_config.TextColumn("抽取值",   width="large"),
                    "原文出处": st.column_config.TextColumn("原文出处", disabled=True, width="medium"),
                    "置信度":   st.column_config.TextColumn("置信度",   disabled=True, width="small"),
                    "校验":     st.column_config.TextColumn("校验",     disabled=True, width="small"),
                    "状态":     st.column_config.TextColumn("状态",     disabled=True, width="small"),
                    "确认":     st.column_config.CheckboxColumn("确认填入"),
                },
                hide_index=True,
                use_container_width=True,
                key="review_editor",
            )

            missing  = [r["字段"] for r in edited if r["状态"] == "未检索到"]
            low_conf = [r["字段"] for r in edited if r["状态"] == "低置信"]
            mismatch = [r["字段"] for r in edited if r["状态"] == "需人工核对"]
            rag_used = [r["字段"] for r in edited if r.get("来源") == "RAG法条"]

            if rag_used:
                st.info(f"🔍 以下字段来自 RAG 知识库：{', '.join(rag_used)}")
            if missing:
                st.warning(
                    f"**未检索到 {len(missing)} 个字段**，需手动补充或保持空白：\n\n- "
                    + "\n- ".join(missing)
                )
            if low_conf:
                st.info(f"**低置信 {len(low_conf)} 个字段**：{', '.join(low_conf)}")
            if mismatch:
                st.error(f"**校验不一致 {len(mismatch)} 个字段**：{', '.join(mismatch)}")

            # ========== Step 4: 回填模板 ==========
            st.divider()
            st.markdown("##### 4. 回填模板")

            col_a, col_b, col_c = st.columns(3)
            also_pdf_2 = col_a.checkbox("同时导出 PDF", value=False)
            interview_demo = col_b.checkbox(
                "询问笔录演示模式", value=False,
                help="勾选后：询问笔录走 AI 问题生成 + RAG 回答",
            )
            custom_prefix = col_c.text_input(
                "输出前缀",
                value=st.session_state.case_dict.get("case_id", "draft"),
            )

            if st.button("确认并回填", type="primary", use_container_width=True):
                mapping = {
                    r["字段"]: r["抽取值"]
                    for r in edited
                    if r["确认"] and r["抽取值"]
                }

                os.makedirs(OUTPUT_DIR, exist_ok=True)
                ts = datetime.now().strftime("%Y%m%d_%H%M%S")
                results = []

                for tpl in tpls:
                    with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp:
                        tmp.write(tpl.getvalue()); tpl_path = tmp.name
                    try:
                        base = os.path.splitext(tpl.name)[0]
                        out_path = os.path.join(
                            OUTPUT_DIR,
                            f"{base}_{custom_prefix}_{ts}.docx",
                        )

                        if "询问笔录" in tpl.name:
                            from generate_interview import generate_interview_docx
                            case_copy = dict(st.session_state.case_dict)
                            case_copy["_generate_answers"] = interview_demo
                            generate_interview_docx(case_copy, out_path)
                        else:
                            doc = Document(tpl_path)
                            fill_paragraph_placeholders(doc, mapping)
                            for table in doc.tables:
                                fill_table_cells(table, mapping)
                            doc.save(out_path)

                        results.append(out_path)
                        if also_pdf_2:
                            pdf_path = docx_to_pdf(out_path)
                            if pdf_path:
                                results.append(pdf_path)
                    except Exception as e:
                        st.error(f"❌ {tpl.name}: {e}")
                    finally:
                        os.remove(tpl_path)

                if results:
                    st.success(f"已生成 {len(results)} 份文件")
                    buf = BytesIO()
                    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
                        for f in results:
                            if os.path.exists(f):
                                zf.write(f, arcname=os.path.basename(f))
                    st.download_button(
                        "📦 打包下载全部",
                        data=buf.getvalue(),
                        file_name=f"文书_{custom_prefix}_{ts}.zip",
                        mime="application/zip",
                        use_container_width=True,
                    )
                else:
                    st.warning("没有生成任何文件")


# ============================================================
# Tab3: 法律问答
# ============================================================
with tab_qa:
    st.markdown("#### 消防法律 RAG 问答")

    try:
        get_deepseek(); get_embedding_config()
        key_ok = True
    except RuntimeError:
        key_ok = False
    if not key_ok:
        st.warning("请先在左侧配置 API Key")

    mode = st.radio("检索模式", ["单文件", "多文件", "全覆盖"],
                    horizontal=True, label_visibility="collapsed")

    from langchain_core.documents import Document
    from langchain_text_splitters import RecursiveCharacterTextSplitter
    from langchain_community.vectorstores import FAISS
    from langchain_openai import ChatOpenAI, OpenAIEmbeddings
    from langchain_core.prompts import ChatPromptTemplate
    from langchain_core.runnables import RunnablePassthrough
    from langchain_core.output_parsers import StrOutputParser
    from pypdf import PdfReader

    def _build_chain_from_files(files_data):
        cfg = get_embedding_config()
        try:
            from openai import OpenAI as _OAI
            _OAI(api_key=cfg["api_key"], base_url=cfg["base_url"]).embeddings.create(
                model=cfg["model"], input="test"
            )
        except Exception as e:
            st.error(f"❌ Embedding Key 无效：{e}")
            st.stop()

        docs = []
        for name, data in files_data:
            with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp:
                tmp.write(data); p = tmp.name
            try:
                reader = PdfReader(p)
                for i, pg in enumerate(reader.pages):
                    t = pg.extract_text() or ""
                    if t.strip():
                        docs.append(Document(page_content=t,
                                             metadata={"source": name, "page": i}))
            finally:
                os.remove(p)
        splitter = RecursiveCharacterTextSplitter(
            chunk_size=chunk_size, chunk_overlap=chunk_overlap,
            separators=["\n\n", "\n", "。", "！", "？", "；", " ", ""],
        )
        chunks = splitter.split_documents(docs)
        emb = OpenAIEmbeddings(
            model=cfg["model"], api_key=cfg["api_key"],
            base_url=cfg["base_url"], check_embedding_ctx_length=False,
        )
        return FAISS.from_documents(chunks, embedding=emb)

    def _build_chain(vs):
        from config import get_deepseek_config
        cfg_ds = get_deepseek_config()
        llm = ChatOpenAI(
            model="deepseek-chat",
            api_key=cfg_ds["api_key"], base_url=cfg_ds["base_url"],
            temperature=0.1, timeout=120, max_retries=2,
        )
        prompt = ChatPromptTemplate.from_messages([
            ("system",
             "你是消防法律专家。请严格根据以下条文回答：\n{context}\n"
             "找不到依据时，直接回答“根据提供的文档无法找到相关法律依据”。"),
            ("human", "{input}"),
        ])
        retriever = vs.as_retriever(search_kwargs={"k": top_k})
        chain = (
            {"context": retriever | (lambda d: "\n\n".join(x.page_content for x in d)),
             "input": RunnablePassthrough()}
            | prompt | llm | StrOutputParser()
        )
        return chain, retriever

    if mode in ("单文件", "多文件"):
        multi = (mode == "多文件")
        files = st.file_uploader(
            "上传法规 PDF" + ("（可多选）" if multi else ""),
            type=["pdf"], accept_multiple_files=multi, key=f"rag_{mode}",
        )
        if files:
            files_list = files if isinstance(files, list) else [files]
            files_data = [(f.name, f.getvalue()) for f in files_list]
            key_fingerprint = f"{ds_key[:10]}_{emb_key[:10]}"
            sig = hashlib.md5(b"".join(f.getvalue() for f in files_list)).hexdigest() \
                  + f"{chunk_size}_{chunk_overlap}_{top_k}_{key_fingerprint}"
            skey = f"rag_{mode}_sig"
            if st.session_state.get(skey) != sig:
                with st.spinner("构建索引中..."):
                    t0 = time.time()
                    vs = _build_chain_from_files(files_data)
                    chain, retriever = _build_chain(vs)
                    bt = time.time() - t0
                st.session_state[f"rag_{mode}_chain"] = chain
                st.session_state[f"rag_{mode}_retriever"] = retriever
                st.session_state[skey] = sig
                st.session_state[f"rag_{mode}_info"] = (len(files_list), bt)

            nf, bt = st.session_state[f"rag_{mode}_info"]
            st.success(f"{nf} 文件 · 建库耗时 {bt:.2f}s")

            q = st.text_input("输入问题", key=f"q_{mode}")
            if q:
                with st.spinner("检索中..."):
                    ans = st.session_state[f"rag_{mode}_chain"].invoke(q)
                    src = st.session_state[f"rag_{mode}_retriever"].invoke(q)
                st.markdown("### 回答")
                st.markdown(ans)
                with st.expander("源片段"):
                    for i, d in enumerate(src, 1):
                        st.markdown(f"**片段 {i}** — {os.path.basename(d.metadata.get('source','?'))}")
                        st.text(d.page_content[:300])
                        st.divider()
    else:
        index_dir = INDEX_DIRS["full"]
        index_file = os.path.join(index_dir, "index.faiss")
        pkl_file = os.path.join(index_dir, "index.pkl")
        if not (os.path.isfile(index_file) and os.path.isfile(pkl_file)):
            st.error(
                "❌ 全覆盖 FAISS 索引不完整。\n\n"
                f"索引目录：{index_dir}\n"
                f"index.faiss：{'存在' if os.path.isfile(index_file) else '缺失'}\n"
                f"index.pkl：{'存在' if os.path.isfile(pkl_file) else '缺失'}\n\n"
                "请先运行 build_index.py（MODE='full'）重新构建索引。"
            )
        else:
            @st.cache_resource(show_spinner="加载索引...")
            def _load_preset(index_path, k, emb_model, emb_url, emb_sig):
                cfg = get_embedding_config()
                emb = OpenAIEmbeddings(
                    model=cfg["model"], api_key=cfg["api_key"],
                    base_url=cfg["base_url"], check_embedding_ctx_length=False,
                )
                vs = FAISS.load_local(index_path, emb, allow_dangerous_deserialization=True)
                return _build_chain(vs)

            cfg_now = get_embedding_config()
            chain, retriever = _load_preset(
                index_dir, top_k,
                cfg_now["model"], cfg_now["base_url"], cfg_now["api_key"][:12],
            )
            q = st.text_input("输入问题", key="q_full",
                              placeholder="例如：浙江省消防条例中关于电动自行车停放的规定是什么？")
            if q:
                with st.spinner("检索中..."):
                    ans = chain.invoke(q)
                    src = retriever.invoke(q)
                st.markdown("### 回答")
                st.markdown(ans)
                with st.expander("源片段"):
                    for i, d in enumerate(src, 1):
                        st.markdown(f"**片段 {i}** — {os.path.basename(d.metadata.get('source','?'))}")
                        st.text(d.page_content[:300])
                        st.divider()