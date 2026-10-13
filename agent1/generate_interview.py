from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import json
import time

from config import ds_client
from rag_module import rag_ask


QUESTION_GEN_PROMPT = """你是一位消防行政执法人员，正在对当事人进行现场询问。
请根据以下案情和检索到的法律条文，生成 5-8 个有针对性的询问问题。

案情：
{case_context}

检索到的法律依据：
{legal_context}

要求：
1. 每个问题必须引用具体法条编号（如"根据《中华人民共和国消防法》第十六条"）
2. 问题应围绕"是否知晓规定""是否履行义务""是否违反法条"展开
3. 不要问纯事实性问题（如"哪一栋""哪一层""什么时间""持续多久"），那些由执法人员现场灵活提问
4. 问题要针对案情中的具体违法违规行为
5. 每个问题独立、简洁明了，不带"问："前缀
6. 返回 JSON：{{"questions": ["问题1", "问题2", ...]}}
"""


def _strip_prefix(q: str, idx: int) -> str:
    """去掉 LLM 可能加的各种前缀（问：、1.、问题：等）"""
    q = q.strip()
    prefixes = [f"{idx}.", f"{idx}、", "问：", "问:", "问题：", "问题:"]
    for p in prefixes:
        if q.startswith(p):
            q = q[len(p):].strip()
    return q


def _generate_ai_answer(question: str) -> str:
    """演示模式：调用 RAG 生成 AI 建议回答"""
    try:
        ans = rag_ask(question)
        # 清理可能的换行
        ans = ans.strip().replace("\n\n", "\n")
        return ans
    except Exception as e:
        return f"（AI 生成失败：{e}）"


def generate_interview_docx(
    case_context: dict,
    output_path: str,
    generate_answers: bool = None,
):
    """
    生成询问提纲 Word 文档

    :param case_context: 案情字典
    :param output_path: 输出路径
    :param generate_answers:
        - None（默认）：从 case_context["_generate_answers"] 读取，没有就 False
        - False：提纲模式 —— AI 只生成问题，答留空（符合执法实务）
        - True ：演示模式 —— AI 生成问题 + RAG 生成回答
    """
    # 参数归一化
    if generate_answers is None:
        generate_answers = bool(case_context.get("_generate_answers", False))

    case_desc = case_context.get("case_context", "")

    # ========== 1. 先用 RAG 检索相关法条 ==========
    print("    [1/4] 检索相关法条...")
    try:
        legal_context = rag_ask(
            f"消防监督执法中，针对以下情况涉及哪些法律条文？请列出条文编号和原文：{case_desc}"
        )
    except Exception as e:
        print(f"    ⚠️ RAG 检索失败：{e}，将使用空法条上下文")
        legal_context = "（未能检索到法条，请基于常识生成通用问题）"

    # ========== 2. 调用 LLM 生成问题 ==========
    print("    [2/4] 调用 LLM 生成询问提纲...")
    resp = ds_client.chat.completions.create(
        model="deepseek-chat",
        messages=[{"role": "user", "content": QUESTION_GEN_PROMPT.format(
            case_context=case_desc,
            legal_context=legal_context,
        )}],
        temperature=0.5,
        response_format={"type": "json_object"},
    )
    questions = json.loads(resp.choices[0].message.content)["questions"]

    # ========== 3. 演示模式：为每个问题生成 AI 回答 ==========
    qa_pairs = []   # [(问题, 回答 or None)]
    if generate_answers:
        print(f"    [3/4] 演示模式：为 {len(questions)} 个问题生成 AI 回答...")
        for i, q in enumerate(questions, 1):
            q_clean = _strip_prefix(q, i)
            print(f"          ({i}/{len(questions)}) {q_clean[:30]}...")
            ans = _generate_ai_answer(q_clean)
            qa_pairs.append((q_clean, ans))
            time.sleep(0.3)   # 避免 API 限流
    else:
        print("    [3/4] 提纲模式：答案留空由办案人员现场记录")
        for i, q in enumerate(questions, 1):
            qa_pairs.append((_strip_prefix(q, i), None))

    # ========== 4. 写入 Word ==========
    print("    [4/4] 写入 Word 文档...")
    doc = Document()
    title = doc.add_heading("询问笔录", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # --- 表头 ---
    doc.add_paragraph(f"询问时间：{case_context.get('time', '____年__月__日____时__分')}")
    doc.add_paragraph(f"询问地点：{case_context.get('location', '____________')}")
    doc.add_paragraph(f"询问人：{case_context.get('interviewer', '____________')}")
    doc.add_paragraph("记录人：____________")
    doc.add_paragraph(f"被询问人：{case_context.get('interviewee', '____________')}")
    doc.add_paragraph("")

    # --- 演示模式的水印提示 ---
    if generate_answers:
        warn = doc.add_paragraph()
        warn.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = warn.add_run("⚠️ 本文档由 AI 生成，仅供内部参考，不得直接作为执法依据 ⚠️")
        run.bold = True
        run.font.size = Pt(10)
        doc.add_paragraph("")

    # --- 第一部分：首问必问事项 ---
    doc.add_paragraph("【首问必问事项】").runs[0].bold = True
    first_items = [
        "你的姓名、性别、年龄、出生日期、身份证件种类及号码、户籍所在地、现住址、联系方式？",
        "你是否曾因消防安全违法行为受过行政处罚？如有，请说明处罚时间、案由及种类。",
    ]
    for q in first_items:
        p_q = doc.add_paragraph()
        p_q.add_run("问：").bold = True
        p_q.add_run(q)
        if generate_answers:
            p_a = doc.add_paragraph()
            p_a.add_run("答：").bold = True
            p_a.add_run("（待被询问人现场陈述）")
        else:
            doc.add_paragraph("答：________________________________________________")
    doc.add_paragraph("")

    # --- 第二部分：权利义务告知 ---
    doc.add_paragraph("【权利义务告知】").runs[0].bold = True
    doc.add_paragraph(
        "问：根据有关法律法规，你享有以下权利：（1）有权要求回避；"
        "（2）有陈述和申辩的权利；（3）对与本案无关的问题，有拒绝回答的权利；"
        "（4）对询问笔录记载有遗漏或差错的，有提出补充或更正的权利。"
        "同时，你应当如实提供证据、证言。你听清楚了吗？"
    )
    if generate_answers:
        p_a = doc.add_paragraph()
        p_a.add_run("答：").bold = True
        p_a.add_run("听清楚了。")
    else:
        doc.add_paragraph("答：________________________________________________")
    doc.add_paragraph("")

    # --- 第三部分：AI 生成的具体问题 ---
    doc.add_paragraph("【案件相关问题】").runs[0].bold = True
    for i, (q_clean, ans) in enumerate(qa_pairs, 1):
        p_q = doc.add_paragraph()
        p_q.add_run(f"{i}. 问：").bold = True
        p_q.add_run(q_clean)

        p_a = doc.add_paragraph()
        p_a.add_run("答：").bold = True
        if ans:
            p_a.add_run(ans)
        else:
            p_a.add_run("________________________________________________")

        doc.add_paragraph("")

    # --- 结尾签名区 ---
    doc.add_paragraph("")
    doc.add_paragraph("以上笔录我看过，与我说的相符。")
    doc.add_paragraph("")
    doc.add_paragraph("被询问人签名：____________   日期：____年__月__日")
    doc.add_paragraph("询问人签名：____________     日期：____年__月__日")
    doc.add_paragraph("记录人签名：____________     日期：____年__月__日")

    doc.save(output_path)
    return output_path