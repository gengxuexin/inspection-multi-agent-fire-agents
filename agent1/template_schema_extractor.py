"""
从任意 docx 模板提取字段清单。
输出：[{key, label, type, where}, ...]

支持：
- 段落占位符  【xxx】 / {{xxx}}
- 段落"标签：___"
- 表格单元格内占位符
- 表格"左标签 + 右空格"
"""
import re
from docx import Document
from document_io import PLACEHOLDER_PATTERN


def _unique_cells(row):
    """同一行中合并单元格会出现重复引用，去重。"""
    seen, result = set(), []
    for cell in row.cells:
        tc_id = id(cell._tc)
        if tc_id not in seen:
            seen.add(tc_id)
            result.append(cell)
    return result


def extract_fields_from_template(docx_path: str) -> list:
    doc = Document(docx_path)
    fields = []
    seen = set()

    def _add(key, label, ftype, where):
        k = key.strip().replace(" ", "").replace("　", "")
        if not k:
            return
        if k in seen:
            return
        seen.add(k)
        fields.append({
            "key": k,
            "label": label or k,
            "type": ftype,
            "where": where,
        })

    # 1. 段落占位符
    for i, p in enumerate(doc.paragraphs):
        for m in PLACEHOLDER_PATTERN.finditer(p.text):
            name = m.group(1) or m.group(2)
            _add(name, name, "placeholder", f"para#{i}")

    # 2. 段落"标签：___"
    for i, p in enumerate(doc.paragraphs):
        m = re.match(
            r"^\s*([\u4e00-\u9fa5A-Za-z0-9_（）()]{2,20})[：:]\s*[_\-—]{2,}",
            p.text,
        )
        if m:
            _add(m.group(1), m.group(1), "label", f"para#{i}")

    # 3. 表格
    for ti, table in enumerate(doc.tables):
        for ri, row in enumerate(table.rows):
            cells = _unique_cells(row)
            for ci, cell in enumerate(cells):
                txt = cell.text.strip()

                # 3a. cell 内占位符
                for m in PLACEHOLDER_PATTERN.finditer(txt):
                    name = m.group(1) or m.group(2)
                    _add(name, name, "placeholder", f"T{ti}R{ri}C{ci}")

                # 3b. 左标签 + 右空格
                if 2 <= len(txt) <= 20 and "【" not in txt and "{{" not in txt:
                    if ci + 1 < len(cells):
                        right = cells[ci + 1].text.strip()
                        if not right:
                            _add(txt.rstrip("：:"), txt, "table_label",
                                 f"T{ti}R{ri}C{ci}")

    return fields


def summarize_fields(fields: list) -> dict:
    """按 type 汇总，便于 UI 展示。"""
    result = {"placeholder": 0, "label": 0, "table_label": 0}
    for f in fields:
        result[f["type"]] = result.get(f["type"], 0) + 1
    return result